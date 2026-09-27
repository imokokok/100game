from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "process"


def make_doc(filename, title, author, intro, sections):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.78)
    sec.left_margin = sec.right_margin = Inches(0.88)
    styles = doc.styles
    for name, size, bold, before, after in [
        ("Normal", 10.5, False, 0, 7),
        ("Title", 27, True, 0, 8),
        ("Subtitle", 11, False, 0, 18),
        ("Heading 1", 17, True, 18, 7),
        ("Heading 2", 12.5, True, 12, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(17, 18, 16)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
    p = doc.add_paragraph(title, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(f"English edition · Written by {author}", style="Subtitle")
    doc.add_paragraph(intro)
    for heading, blocks in sections:
        doc.add_paragraph(heading, style="Heading 1")
        for kind, text in blocks:
            if kind == "h2":
                doc.add_paragraph(text, style="Heading 2")
            elif kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)
    doc.core_properties.title = title
    doc.core_properties.author = author
    doc.core_properties.subject = "HOW 100 PEOPLE CALL A GAME process archive"
    doc.save(OUT / filename)


make_doc(
    "week1-game-narrative-design-en.docx",
    "Game Narrative Design 1.0",
    "咸鱼",
    "This document establishes the narrative foundation for the island town, the seven-day application process and the parallel lives of A and B. Its tone begins with warmth, humour and everyday absurdity, then allows emotional weight to emerge from choices the player has already made.",
    [
        ("1. Setting: an island town that is not open to the public", [
            ("p", "The story takes place in a fictional island town with a strong local culture and an unbroken pattern of traditional life. Its old streets hold handmade goods, family-run shops, morning markets and night markets. Life moves slowly, and people know one another well enough that passing someone in the street usually means stopping to say hello."),
            ("p", "The town maintains a deliberate distance from the outside world. It is not completely closed, but newcomers must live by rules that the local community understands better than visitors do. The island university, provisionally called Island University, leads in a small number of specialist fields and attracts applicants from very different backgrounds. A and B are both applying. Before they can enrol, each must complete a seven-day preliminary stay and receive recognition from at least twelve permanent residents."),
            ("p", "The town draws on the everyday texture of Tainan: old streets, temple forecourts, markets, lanes, sea wind, cicadas and long summer evenings. Magical realism appears through ordinary encounters rather than spectacle. At an unremarkable corner, a girl may be reading tarot. A strange flute may sound in an alley after midnight. A restaurant menu may include a dish called 'What Time Do We Get Off Work?'"),
            ("h2", "Residency and study permit"),
            ("p", "The university does not recruit publicly. Applicants first enter the island as temporary residents, live there for seven days, keep a record of daily life and obtain at least twelve statements from permanent residents confirming that they know the applicant. A relationship may be close or may consist only of repeated encounters. A pleasant conversation does not automatically earn recognition, and residents may refuse. The system expresses a simple idea: a form cannot place someone inside a community. A person has to be seen there in specific, lived ways."),
        ]),
        ("2. Two parallel routes across seven days", [
            ("p", "The game begins on the morning of Day 1 and ends after the review on Day 7. A and B live through the same seven days in the same town. The player may choose whose route to begin and when to switch. Their time runs in parallel: while A is doing one thing, B is living through another."),
            ("bullet", "Days 1–3: establish rhythm. A wanders, follows chance meetings and lets a day expand around them. B plans routes, checks opening hours and records schedules."),
            ("bullet", "Days 4–5: choices accumulate. Their circles of recognition begin to take shape, and pressure around money and time becomes visible. B may need paid work to cover food and accommodation. A may realise that too much of her time has gone to things that do not support her application."),
            ("bullet", "Day 6: the last window. The player must decide which relationships can still deepen and which confirmations remain possible."),
            ("bullet", "Day 7: review. The outcome reflects the previous six days. There is no simple good or bad ending, only what the player left in the town and how many people truly saw them."),
            ("h2", "Where the routes cross"),
            ("p", "A and B may pass one another several times. A may eat at the restaurant where B works. B may notice A talking with a painter. They may know that the other is also an applicant without ever knowing her well. The player sees two routes crossing within one space, while neither character has access to the whole picture. That difference in knowledge is part of the subject of the game."),
        ]),
        ("3. Character design", [
            ("h2", "A: act first, reflect later"),
            ("p", "A comes from a financially comfortable family. When she encounters a new place, a stranger or an unexpected invitation, she is more likely to enter first and adjust afterwards. Her ease comes from knowing that a mistake is unlikely to ruin everything. She applies because she genuinely loves a specialist subject taught exceptionally well on the island. She has long assumed she is talented in it, but gradually discovers that what she loves may be precisely what does not come naturally to her."),
            ("p", "A cooks, daydreams, walks without a fixed route and talks to people. Temporary social encounters can enlarge the possibilities of a day. Buying someone a meal may lead to an introduction. Her difficulty is that she is used to solving problems with money and has little practice with problems that money cannot solve. The discovery that love and talent are not the same threatens her sense of who she is."),
            ("h2", "B: reflect first, then act"),
            ("p", "B comes from a family with fewer financial choices. She checks location, cost, opening hours, transport and people's schedules before deciding where to go. Her caution comes from experience: a wrong decision can carry serious consequences. Her reason for applying is practical. The university works with employers, and she hopes that studying there will lead to stable work. She may not describe the subject as a passion, but she knows what she needs."),
            ("p", "B works in exchange for food and accommodation, records schedules and routes in a notebook, and doubts herself late at night. Many of her relationships begin through work: colleagues, regular customers and children she meets while serving them. She has learned to fill gaps in resources with preparation and effort, but not every problem yields to effort. In a quiet moment, she may remember being dismissed as a child. In another, she may realise that she has already travelled much farther than she allows herself to believe."),
            ("h2", "Seeing one another"),
            ("p", "A may admire the way B remains warm and capable with every customer. B may admire A's fluency across unfamiliar subjects and the small kindnesses she offers without making them dramatic. They do not need to become friends. They only need the chance to notice something in the other that they themselves lack. The game's idea of understanding is not pity or rescue. It is the act of seeing another person."),
        ]),
        ("4. Narrative approach and tone", [
            ("p", "The opening should be light, playful and slightly absurd. Its later weight grows out of the player's earlier decisions. The writing moves towards magical realism through specific places, routines and personal experience. It should not use philosophical dialogue to explain the meaning of a scene on the player's behalf."),
            ("p", "The everyday reference is the warmth, humour and close observation of A Taiwanese Tale of Two Cities-style domestic storytelling: old streets, markets, temple forecourts and summer nights. A and B are ordinary people rather than geniuses or chosen heroes. They are trying to live well inside different conditions."),
            ("bullet", "Accumulated detail: growth is woven from small observations rather than driven only by major events."),
            ("bullet", "Interlaced routes: both women occupy the same town and occasionally pass one another; the player sees what neither individual character can."),
            ("bullet", "Light-touch interaction: photographs, drawings, collage and sound carry chapter transitions without interrupting the story with heavy controls."),
        ]),
        ("5. One-sentence summary", [
            ("p", "Two ordinary people spend the same seven days in an island town that is closed to the public, each moving through it in a completely different way. The game asks how we recognise one another when our resources and histories differ, and whether complete understanding is necessary before we acknowledge that another person is here."),
        ]),
    ],
)


def screenplay_pdf(filename, title_text, author, scenes):
    path = OUT / filename
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ScreenTitle", parent=styles["Title"], fontName="Courier-Bold", fontSize=22, leading=26, textColor=HexColor("#111210"), spaceAfter=8)
    meta_style = ParagraphStyle("ScreenMeta", parent=styles["Normal"], fontName="Courier", fontSize=9, leading=13, textColor=HexColor("#686a66"), spaceAfter=16)
    slug_style = ParagraphStyle("Slug", parent=styles["Heading2"], fontName="Courier-Bold", fontSize=11, leading=14, textColor=HexColor("#111210"), spaceBefore=12, spaceAfter=6)
    action_style = ParagraphStyle("Action", parent=styles["BodyText"], fontName="Courier", fontSize=9.5, leading=14, textColor=HexColor("#111210"), spaceAfter=6)
    character_style = ParagraphStyle("Character", parent=styles["BodyText"], fontName="Courier-Bold", fontSize=9.5, leading=12, leftIndent=58*mm, spaceBefore=7, spaceAfter=1)
    dialogue_style = ParagraphStyle("Dialogue", parent=styles["BodyText"], fontName="Courier", fontSize=9.5, leading=13, leftIndent=40*mm, rightIndent=40*mm, spaceAfter=5)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22*mm, rightMargin=22*mm, topMargin=18*mm, bottomMargin=18*mm, title=title_text, author=author)
    story = [Paragraph(title_text, title_style), Paragraph(f"SCREENPLAY · English edition · Written by {author}", meta_style), Paragraph("FADE IN:", slug_style)]
    for heading, lines in scenes:
        story.append(Paragraph(heading, slug_style))
        for kind, text in lines:
            story.append(Paragraph(text.replace("&", "&amp;"), character_style if kind == "c" else dialogue_style if kind == "d" else action_style))
    story.append(Paragraph("CUT TO BLACK.", slug_style))
    doc.build(story)


screenplay_pdf(
    "week2-solmere-script-a-en.pdf",
    "SOLMERE · A BEFORE SOLMERE",
    "HuieChen 陈慧娥",
    [
        ("1. INT. A'S BEDROOM / DINING ROOM — MORNING", [
            ("a", "The curtains are closed. The phone alarm has already rung once. A lies awake in bed. Her mother knocks twice."),
            ("c", "MOTHER (O.S.)"), ("d", "Aren't you going to class today?"),
            ("c", "A"), ("d", "I don't feel like it."),
            ("a", "Silence outside the door. Then her mother's footsteps move away. A turns over, then back again. She checks her phone. No messages from either parent."),
            ("a", "Later, her mother is eating breakfast. A enters in school uniform and grabs a slice of bread."),
            ("c", "MOTHER"), ("d", "I thought you weren't going."),
            ("c", "A"), ("d", "Changed my mind."),
            ("c", "MOTHER"), ("d", "Oh."),
            ("c", "A"), ("d", "Why do you always say 'oh'?"),
            ("c", "MOTHER"), ("d", "What am I supposed to say?"),
            ("c", "A"), ("d", "…Nothing."),
            ("a", "A leaves with the bread. Her mother continues breakfast."),
        ]),
        ("2. INT. LIVING ROOM — EVENING", [
            ("a", "A throws her bag onto the sofa. Her father sits on the carpet taking apart a broken speaker. A hands him a sheet of paper."),
            ("c", "FATHER"), ("d", "First place?"), ("c", "A"), ("d", "Mm."), ("c", "FATHER"), ("d", "Congratulations."),
            ("a", "He returns to the speaker. A does not take the paper."),
            ("c", "FATHER"), ("d", "What's wrong?"), ("c", "A"), ("d", "That's it?"),
            ("c", "FATHER"), ("d", "It's impressive. Should I congratulate you again?"),
            ("a", "A cannot help laughing. She takes the paper and starts towards her room."),
            ("c", "FATHER"), ("d", "What do you want for dinner?"), ("c", "A"), ("d", "Anything."),
            ("c", "FATHER"), ("d", "That answer doesn't work. This one you have to choose."),
            ("a", "A stops and looks back."),
        ]),
        ("3. INT. KITCHEN — NIGHT", [
            ("a", "A sits beside a cup of ice that has half melted. Her mother washes glasses at the sink."),
            ("c", "A"), ("d", "I'm quitting."), ("c", "MOTHER"), ("d", "All right."),
            ("c", "A"), ("d", "That's all? Aren't you going to ask why?"),
            ("c", "MOTHER"), ("d", "Do you want to tell me?"), ("c", "A"), ("d", "No."), ("c", "MOTHER"), ("d", "Then I won't ask."),
            ("c", "A"), ("d", "We already paid for it. The money's wasted."),
            ("c", "MOTHER"), ("d", "It is already wasted. If you keep going now, the money still won't come back."),
            ("c", "A"), ("d", "Why are you both like this?"), ("c", "MOTHER"), ("d", "Like what?"),
            ("a", "A cannot explain. Her mother opens the fridge and takes out fruit."),
            ("c", "MOTHER"), ("d", "Want some?"), ("c", "A"), ("d", "Yes."),
        ]),
        ("4. MONTAGE — AN UNPLANNED SUMMER", [
            ("a", "A and friends run into the sea. On the back seat of a bus they decide, at the last moment, to visit another city. Outside a convenience store after midnight, A laughs until she crouches on the pavement. In a kitchen, she serves something badly cooked."),
            ("c", "FRIEND"), ("d", "What is that?"), ("c", "A"), ("d", "No idea."),
            ("a", "Everyone laughs. Afternoons pass through films, roadside drinks and waking after the sun is already high. A receives a message: 'Out tonight?' She replies, 'I'm in.'"),
            ("a", "At sunset on the steps by the sea, the group discusses tomorrow."),
            ("c", "FRIEND 1"), ("d", "Decide after we wake up?"), ("c", "A"), ("d", "Fine."),
            ("c", "A"), ("d", "What date is it?"), ("c", "FRIEND 1"), ("d", "The twenty-seventh."),
            ("c", "A"), ("d", "Already? What did we even do this month?"),
            ("a", "They try to remember and laugh when one memory turns out to belong to the previous month. A laughs too, then checks the date on the phone again."),
        ]),
        ("5. INT. A'S BEDROOM — NIGHT / NEXT MORNING", [
            ("a", "A fills a blank digital calendar: 09:00 WORK. 11:00 READ. 14:00 PROJECT. 17:00 GYM. 20:00 FILM. The empty page quickly becomes full. She leans back, satisfied."),
            ("a", "The next morning, the alarm sounds at 08:30. A silences it. Nobody knocks. Nobody messages. Nobody knows she is still in bed. At 08:47 she stares at the time, slowly sits up and remains on the edge of the bed with her eyes closed. Eventually she stands."),
        ]),
        ("6. INT. SMALL GAME STUDIO — 03:14 / NEXT AFTERNOON", [
            ("a", "Coffee cups, cables and takeaway boxes fill the studio. Someone sleeps across a desk. A programmer compares two versions while A stands behind him."),
            ("c", "PROGRAMMER"), ("d", "Left or right?"), ("c", "A"), ("d", "Right."),
            ("c", "PROGRAMMER"), ("d", "If we change to the right one now, it may break tomorrow."), ("c", "A"), ("d", "Right."),
            ("c", "PROGRAMMER"), ("d", "Certain?"), ("c", "A"), ("d", "Certain."),
            ("c", "PROGRAMMER"), ("d", "All right. Your call."),
            ("a", "The next afternoon, the team watches a test. It passes. Then passes again. The room erupts."),
            ("c", "TEAM MEMBER"), ("d", "Who said to keep the right version?"), ("c", "PROGRAMMER"), ("d", "Her."),
            ("c", "A"), ("d", "Me."),
            ("a", "A laughs, genuinely delighted. Much later, the team eats cold takeaway. The programmer lies on the floor."),
            ("c", "PROGRAMMER"), ("d", "What time tomorrow?"), ("c", "A"), ("d", "Nine."),
            ("c", "PROGRAMMER"), ("d", "You push too hard."), ("c", "A"), ("d", "You're not coming?"),
            ("c", "PROGRAMMER"), ("d", "I'm coming."),
            ("a", "They both smile. A nearly falls asleep to keyboards, discussion and machines, still smiling."),
        ]),
        ("7. INT. GAME STUDIO / A'S WORKSPACE — NIGHT AND NEXT MORNING", [
            ("a", "The project ends. Computers switch off, cables are unplugged and the wall schedule comes down. A uploads the final file. The progress bar reaches 100 percent. Everyone cheers; a drink sprays foam across the room."),
            ("a", "The next morning A wakes naturally. There is no alarm and no work message. The final project chat reads: 'Thank you for everything!!!' She smiles, then finds the whiteboard still carrying the last days of the project: TEST. BUILD. FINAL. DEADLINE."),
            ("a", "A erases each word. When the board is completely empty, she lifts a marker. The tip touches the surface. She waits for a long time and writes nothing."),
        ]),
    ],
)


screenplay_pdf(
    "week2-solmere-script-b-en.pdf",
    "SOLMERE · B BEFORE SOLMERE",
    "HuieChen 陈慧娥",
    [
        ("1. INT. B'S FAMILY HOME, DINING ROOM — NIGHT", [
            ("a", "B sits with her parents at a table more generous than usual. Her father pours her a drink."),
            ("c", "FATHER"), ("d", "Order anything you want today."), ("c", "B"), ("d", "We already ordered."),
            ("c", "FATHER"), ("d", "Then order more."), ("c", "B"), ("d", "You're paying?"),
            ("c", "FATHER"), ("d", "I am today."), ("c", "MOTHER"), ("d", "When don't you pay?"),
            ("a", "They laugh. B likes the food. Her father checks the price on the menu."),
            ("c", "FATHER"), ("d", "Once you have a good job, you can eat somewhere like this whenever you want."),
            ("c", "MOTHER"), ("d", "Which is why you still need to study properly."),
            ("c", "B"), ("d", "How did dinner turn into this again?"),
            ("c", "FATHER"), ("d", "We're not asking anything impossible. When is your exam?"),
        ]),
        ("2. INT. LIVING ROOM — NIGHT", [
            ("a", "School and course materials cover the coffee table. B sits on the carpet while her father holds one brochure."),
            ("c", "FATHER"), ("d", "This one will make it easier to find work."), ("c", "B"), ("d", "I know."),
            ("c", "FATHER"), ("d", "Don't only think about whether you like it. Our family cannot afford too many wrong turns. Decide early, before you have nowhere left to go."),
            ("a", "B opens her notebook and writes: TUITION. EMPLOYMENT. CITY. TIME. PLAN B. After her parents go to bed, she ranks every option. The brochure she studied longest does not end up on top."),
        ]),
        ("3. INT. B'S BEDROOM — LATE NIGHT", [
            ("a", "Music software fills the screen. B records a cup, then keys, then a radiator. She drags the sounds into a track. This time it works, and she smiles. Her mother opens the door."),
            ("c", "MOTHER"), ("d", "Still awake? Making music again?"), ("c", "B"), ("d", "Yes. Everything else is finished."),
            ("c", "B"), ("d", "Mum, do you think it sounds good?"),
            ("a", "Her mother listens carefully."), ("c", "MOTHER"), ("d", "It does."),
            ("c", "MOTHER"), ("d", "What can you do with it later?"), ("c", "B"), ("d", "Listen to it."),
            ("c", "MOTHER"), ("d", "I know. I mean afterwards."), ("c", "B"), ("d", "I'll think about that later."),
            ("a", "After the door closes, B plays the passage from the beginning. She does not edit it again. She only listens."),
        ]),
        ("4. INT. KITCHEN — NIGHT / DAY", [
            ("c", "B"), ("d", "Mum, have you ever wanted to do something you liked, just for yourself? Didn't you want to learn to dance?"),
            ("c", "MOTHER"), ("d", "That was years ago. I don't have time."),
            ("a", "B produces her phone. She has already found nearby classes and checked her mother's Wednesday evening."),
            ("c", "B"), ("d", "This one is twenty minutes away. This one is farther, but—"),
            ("a", "Her mother does not take the phone. B stops."),
            ("c", "B"), ("d", "Am I arranging your life again?"), ("c", "MOTHER"), ("d", "You finally noticed."),
            ("c", "B"), ("d", "Then you choose."),
            ("a", "Her mother accepts the phone and promises to look. In a matching daytime scene later, B notices the course flyer weighted beneath a bag of vegetables. She says nothing."),
        ]),
        ("5. INT. PRODUCTION STUDIO — NIGHT", [
            ("a", "Several problems arrive at once: tomorrow's location is lost, equipment has not arrived, and one person cannot attend. B edits the schedule while everyone speaks."),
            ("c", "B"), ("d", "Move the location to Thursday. If the equipment is not here by eleven, use the backup. Move his section later; I'll connect the earlier material first."),
            ("c", "TEAM MEMBER"), ("d", "Can that really fit?"), ("c", "B"), ("d", "Yes."),
            ("a", "The studio starts moving again. A colleague says it is lucky she is here. B asks for dinner in return, then specifies an expensive one. They laugh."),
            ("a", "Late at night, EXPORT COMPLETE appears. B checks the file, closes it, reopens it and checks again. Her mother texts: 'Still at work?' then, 'Don't take every problem onto yourself.' B stares at the message and continues checking."),
        ]),
        ("6. INT. FAMILY DINING ROOM — EVENING", [
            ("a", "Her father passes B a job listing. The pay and benefits look strong. Lower on the page appears SOLMERE — CREATIVE INDUSTRIES — MEDIA / PRODUCTION. B scrolls through a record shop, theatre, small venues, chess club and art spaces."),
            ("c", "MOTHER"), ("d", "What do you think? Is it stable?"), ("c", "B"), ("d", "There seem to be plenty of jobs. It should be stable."),
            ("a", "A photograph of the record shop fills the screen. When her mother leans closer and asks what she is looking at, B scrolls back to the vacancy."),
            ("c", "B"), ("d", "Work."),
        ]),
        ("7. INT. B'S BEDROOM — NIGHT", [
            ("a", "An open suitcase lies on the bed. B's notebook lists transport, accommodation, budget, work, seven days, twelve recognitions and the return journey. Her mother helps fold clothes."),
            ("a", "B notices a small recorder on the desk and adds it to the suitcase."),
            ("c", "MOTHER"), ("d", "For work?"), ("c", "B"), ("d", "No."), ("c", "MOTHER"), ("d", "Then why bring it?"), ("c", "B"), ("d", "For fun."),
            ("c", "MOTHER"), ("d", "Don't have so much fun that you forget the important part."), ("c", "B"), ("d", "I know."),
            ("a", "After her mother leaves, B scrolls past jobs to the record shop, theatre, chess club, community centre and people in the street. On the notebook's last line she writes: WHEN I COME BACK— She waits, crosses it out and closes the book."),
        ]),
    ],
)


screenplay_pdf(
    "week2-epilogue-en.pdf",
    "SOLMERE · EPILOGUE",
    "HuieChen 陈慧娥",
    [
        ("EXT. REVIEW BUILDING — DAY", [
            ("a", "A opens the door and walks down the steps with a folder in her hand. She looks at it and cannot stop herself from smiling."),
            ("a", "A usually smiles quickly and completely. This time the smile is only a faint line. She has not yet understood how deep the happiness goes. She walks several steps, still smiling, then seems to realise what has happened and lets out one small laugh."),
            ("a", "The door opens again. B steps outside with one hand resting on the bag that holds her documents. Through the fabric, she feels the hard edge of the folder, the thickness of stacked paper and one corner rising beside the zip."),
            ("a", "She knows the process is finished and the material is there. Still, she checks again. The edge. The paper. The raised corner. All real. She almost finds the habit ridiculous, but touches the bag once more before looking up."),
        ]),
        ("INT. RECORD SHOP — LATER", [
            ("a", "The record-shop lights are on. CDs pass one by one through B's fingers. She studies the paper sleeves and labels without knowing what she wants. Perhaps she needs the physical feeling of music. Perhaps she is waiting for one title to suit her. The possibilities remain mixed, and she does not try to separate them."),
            ("a", "A sits at the listening table before a portable CD player. She opens a case, removes the disc, closes the lid and presses PLAY."),
            ("a", "CLICK."),
            ("a", "B looks up. She sees the paper sleeve on the table, then her own name printed on the label."),
            ("a", "A's fingers stop tapping the rhythm. She opens her eyes. Across the room, each woman finds the other person's eyes."),
            ("a", "A warm wind enters from the street. B smiles first. A looks down at the CD in her hand, then back up, and smiles too. Neither knows why. The wind enters again. Perhaps the summer air is simply too warm."),
        ]),
    ],
)


make_doc(
    "week1-town-flow-draft-en.docx",
    "Town Map and Game Flow Draft",
    "咸蛋黄",
    "A working design for the town map, the activities attached to each place, and the alternating routes of protagonists A and B from the prologue to the final day.",
    [
        ("Town map and activity spaces", [
            ("h2", "Restaurant or milk-tea shop · B main route"),
            ("p", "Supports restaurant management and the public recipe-book activity. A can visit and react to recipes left by other players, but cannot contribute without the recipe notebook. A milk-tea shop may be a simpler production alternative."),
            ("h2", "Bookshop or post office · B main route"),
            ("p", "Supports commissioned writing and collage poetry. A can visit and react to other players' work, but cannot take commissions without the pen item."),
            ("h2", "Tarot stall · side route"),
            ("p", "Hosts card reasoning and tarot. Either protagonist can enter at any stage."),
            ("h2", "Chess stall · side route"),
            ("p", "Hosts board-game activities. Either protagonist can enter at any stage."),
            ("h2", "Observatory · side route"),
            ("p", "Supports quiet observation and the world-landscape activity. Either protagonist can enter at any stage."),
            ("h2", "Gallery or science museum · A main route"),
            ("p", "Hosts the three-dimensional optical-illusion activity. B can visit as an exhibition space, but without the phone item she cannot trigger its puzzle layer."),
            ("h2", "Record shop · A main route"),
            ("p", "Hosts field recording, composition and music-video production. B can listen and react to music left by other players, but cannot submit work without headphones."),
            ("h2", "Resolving misunderstandings · shared main route"),
            ("p", "The activity is part of both routes. Its detailed interaction design remains open."),
        ]),
        ("Prologue", [
            ("p", "The opening appears as a blank picture book with four panels arranged two by two. Only the upper-left image is available at first. The player opens each panel in sequence and learns the main interaction language through the characters' journeys."),
            ("h2", "A wakes at home"),
            ("p", "A wakes on her bed in a relaxed mood. She treats the seven-day residency challenge more like a journey than an examination. She packs her phone and headphones and leaves. The scene can establish a little of her personality and circumstances before returning to the picture book."),
            ("h2", "A on the train"),
            ("p", "A listens to music by the window. The scene teaches field recording for the city-sampling activity. The phone then teaches photography and the memory system: selecting an older image reveals part of A's past, while selecting a newly taken image can return her to the photographed location to search for clues."),
            ("h2", "B at her desk"),
            ("p", "B records an outline of the town in a notebook, introducing the map and her careful preparation. Tabs divide the notebook into recipes, town information, accounts and other areas. Her favourite drink recipe opens the milk-tea tutorial. She then turns to her journal materials and learns the basic collage-poetry interaction."),
            ("h2", "B on the ferry"),
            ("p", "B sits facing the window. The player drags the moon down outside the glass. As the moon sets, the sun rises and the next chapter begins."),
        ]),
        ("Weekly chapter flow", [
            ("h2", "Monday · B viewpoint"),
            ("p", "A and B meet at the entrance to town. A starts the conversation and shares one earbud. B responds by showing A her notebook. From this point, A can use a map during her route and B can collect recordings to send to A. The two then separate. B's main objective uses information in the notebook to introduce the town and its residents."),
            ("h2", "Tuesday · A viewpoint"),
            ("p", "The main route builds a puzzle around A's photograph-based return ability, deepening her knowledge of the town and its residents."),
            ("h2", "Wednesday · B viewpoint"),
            ("p", "B enters the bookshop or post office. This chapter may exchange places with Friday."),
            ("h2", "Thursday · A viewpoint"),
            ("p", "A enters the record shop. This chapter may exchange places with Saturday."),
            ("h2", "Friday · B viewpoint"),
            ("p", "B enters the restaurant or milk-tea shop. This chapter may exchange places with Wednesday."),
            ("h2", "Saturday · A viewpoint"),
            ("p", "A enters the gallery or science museum. This chapter may exchange places with Thursday."),
            ("h2", "Sunday · A and B"),
            ("p", "The routes converge in an activity about resolving misunderstanding. The exact design remains to be developed."),
        ]),
    ],
)


make_doc(
    "week3-text-design-example-en.docx",
    "Dialogue Design Sample: Shiyongqi",
    "熊韬炀",
    "This sample shows how Shiyongqi's restaurant dialogue moves from a first meeting to ordinary conversation, deeper trust and recognition earned through play. Her voice is brisk, teasing and practical, but never cold.",
    [
        ("First meeting", [
            ("p", "SHIYONGQI: You're the new one? Come over here. Wash your hands, put on gloves and stand on this side. Don't rush in yet. Let me show you how it works first."),
            ("h2", "Choice 1 · 'No need. I'm fairly used to this.'"),
            ("p", "SHIYONGQI: That's a confident answer. All right, sort this bowl of vegetables. Yellow leaves and tough roots come off. Do it properly. I don't do careless work here. Finish that, and then we'll talk about the next step."),
            ("p", "After the preparation task: SHIYONGQI: Your hands know what they're doing. Maybe you really are used to this. So, are you working with me from today, or do you want to look around first?"),
            ("h2", "Choice 2 · 'All right. I'll watch carefully first.'"),
            ("p", "After the demonstration: SHIYONGQI: That's about it. Want to try it yourself?"),
            ("p", "Accept: after the short preparation task, SHIYONGQI says, 'Not bad. So, are you working with me from now on?'"),
            ("p", "Decline: SHIYONGQI says, 'All right. I can't keep you here if you don't want to stay. Come back when you've decided.'"),
            ("p", "Further choices can lead into another restaurant activity or allow the player to leave and explore before returning."),
        ]),
        ("Later interaction: everyday conversation", [
            ("p", "SHIYONGQI: You're back. Are you here to work today, or have you come to borrow my air-conditioning?"),
            ("h2", "Business"),
            ("p", "PLAYER: Busy today?"),
            ("p", "SHIYONGQI: Busy. That's when it gets interesting. When it's quiet, I end up talking to myself like an idiot. Sometimes I talk to the wok. It never answers. Very dull company."),
            ("h2", "Health"),
            ("p", "PLAYER: You stand all day. Doesn't it exhaust you?"),
            ("p", "SHIYONGQI: My legs barely belong to me by closing time, so yes. My back is still hanging on. Honestly, sitting down can be worse. Once I'm down, I can't get up. If you see me sitting there staring into space, pull me up or I'll stay there until closing."),
            ("h2", "Recommendation"),
            ("p", "PLAYER: What do you recommend?"),
            ("p", "SHIYONGQI: Bullfrog and fish hotpot. Properly spicy. You'll remember me for three days and wonder whether you've burnt your tongue off. I can make it mild if you don't eat chilli, but I'll respect you a tiny bit less. I'm joking. Mostly."),
        ]),
        ("Deeper conversation · unlocks after repeated interaction", [
            ("p", "SHIYONGQI: Sit down. Let's talk. Do you genuinely want to stay here, or do you simply have nowhere else to go? Answer honestly. I won't laugh at either one."),
            ("h2", "1 · 'I want to stay here.'"),
            ("p", "SHIYONGQI: Good. When opening this place was at its worst, someone was willing to sit here and let me finish talking. No advice, no argument, just listening. If you ever decide not to come back, tell me first. Don't disappear."),
            ("h2", "2 · 'A bit of both.'"),
            ("p", "SHIYONGQI: Honest. I like honest people. You can take your time deciding what comes next. You don't have to hand in an answer to your life today."),
            ("h2", "3 · 'I actually like cooking.'"),
            ("p", "SHIYONGQI: Liking it is enough. If you care about it, I'll teach you a few extra things. Have you visited the produce stall? If you know the owner, you may find a surprise. She keeps a few strange ingredients off the signboard. Odd flavours. I like them."),
            ("p", "This stage can naturally lead into Shiyongqi's relationships with Mingming and BEETMAN."),
        ]),
        ("Earning recognition through play", [
            ("h2", "1 · Complete a full rush"),
            ("p", "SHIYONGQI: Orders will pile up tonight. I won't have time to teach you one step at a time. Don't keep asking whether something is all right. Start. If it's wrong, do it again. If you're slow, move faster. We get through this together."),
            ("p", "After service: SHIYONGQI: You made it through, and your hands didn't shake badly enough to throw the wok. You pass. Next time we're busy, start without asking me."),
            ("h2", "2 · Leave one original dish"),
            ("p", "SHIYONGQI: Use what's left and make something of your own. If it works, write it in my book. Once it's there, don't tear the page out. I hold grudges."),
            ("p", "After cooking: SHIYONGQI: Fine. The flavour works. At least it isn't a culinary disaster. Here's the book. Write it down yourself. Your handwriting can be ugly as long as I can read it. The dish matters more, doesn't it?"),
            ("h2", "3 · Bring back BEETMAN's special ingredient"),
            ("p", "SHIYONGQI: Go to the produce stall and use my name. If BEETMAN gives you something strange, bring it back. I'll taste it. If it survives that, it counts in your favour."),
            ("p", "After the ingredient is cooked: SHIYONGQI: Let me see. That flavour really is strange. All right, I'll take it. The trip wasn't wasted."),
            ("h2", "Recognition complete"),
            ("p", "SHIYONGQI: All right. I recognise you. From now on, part of this kitchen is yours. Don't get carried away. One third at most. I still have to watch the rest or you'll burn the place down. I once dreamed I turned into a bullfrog and threw myself into the pot. I woke up and stared at the ceiling for ages, genuinely wondering whether something was wrong with me. If you ever have a dream that ridiculous, come and tell me. It feels like the beginning of some paranormal case. Honestly, that could be interesting."),
        ]),
        ("Triggered dialogue", [
            ("bullet", "Mention Shiyongqi at the produce stall: BEETMAN offers an off-menu suggestion for an unusual flavour."),
            ("bullet", "Mention Mingming at the restaurant: Shiyongqi falls briefly silent, then says, 'She saw me at my worst. Having someone stay beside you at a time like that is… good.' She then changes the subject."),
        ]),
    ],
)


def build_text_design_pdf():
    path = OUT / "week3-text-design-2-en.pdf"
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=HexColor("#111210"), spaceAfter=10)
    meta = ParagraphStyle("Meta", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=14, textColor=HexColor("#686a66"), spaceAfter=14)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=HexColor("#e6543f"), spaceBefore=8, spaceAfter=7)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=HexColor("#111210"), spaceBefore=6, spaceAfter=4)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=14, textColor=HexColor("#111210"), alignment=TA_LEFT, spaceAfter=6)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=16*mm, rightMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm, title="Narrative Design: Maya Bookshop Interaction", author="山梨之心")
    story = [Paragraph("Narrative Design: Maya Bookshop Interaction", title), Paragraph("English edition · Written by 山梨之心", meta)]
    story += [Paragraph("Framework", h1), Paragraph("B follows the same overall interaction structure as A, but the wording must never be identical. Small talk, relationship information and tasks received from the same NPC should reflect what each protagonist already knows and how each one approaches the encounter.", body)]
    story += [Paragraph("First meeting", h1), Paragraph("The first meeting should account for the NPC's schedule and location. It supports three text paths: ordinary conversation that may connect to information about related residents; a relationship-specific interaction; and task recognition.", body)]
    data = [
        [Paragraph("Time and place", h2), Paragraph("Player approach", h2), Paragraph("Maya response and continuation", h2)],
        [Paragraph("08:00–08:30<br/>Outside the bookshop", body), Paragraph("A · Yes. May I come in and have a look?", body), Paragraph("MAYA: I'd be delighted. You're my first customer today.", body)],
        [Paragraph("08:00–08:30<br/>Outside the bookshop", body), Paragraph("B · No. It looks as if it may rain. Do the plants still need watering?", body), Paragraph("MAYA: From what I've seen over the years, it probably won't be heavy today. The player may continue into rain-related small talk.", body)],
        [Paragraph("08:00–08:30<br/>Outside the bookshop", body), Paragraph("C · No, but you know why I'm here, don't you?", body), Paragraph("MAYA: I've met plenty of applicants. You're doing fairly well. The player can continue into task-related dialogue.", body)],
    ]
    table = Table(data, colWidths=[48*mm, 90*mm, 125*mm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),HexColor("#111210")),("TEXTCOLOR",(0,0),(-1,0),HexColor("#ffffff")),("VALIGN",(0,0),(-1,-1),"TOP"),("GRID",(0,0),(-1,-1),0.5,HexColor("#d5d6d2")),("BACKGROUND",(0,1),(-1,-1),HexColor("#f8f7f2")),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
    story += [table, Spacer(1, 8*mm), Paragraph("Rain conversation", h1), Paragraph("Maya says that whenever it rains she comes outside to watch the sky and the rain. Plenty of books have been written about rain, but none compares with seeing it first-hand. Before the rain arrives, the clouds darken and daylight suddenly falls away. The first drops are thin and scattered, tapping on the leaves before falling in a dense curtain.", body), Paragraph("If the player checks the sky and says that the rain looks close, Maya answers that it may begin at any moment and hurries inside. In another branch, Maya remembers delivering books to the mountain district with Mossner: mountain roads are hardest in rain, when the earth turns soft and mud covers everything. Carrying a full load of books through it was no easy task.", body)]
    story += [PageBreak(), Paragraph("Inside the bookshop · 08:30–18:00", title), Paragraph("MAYA: Hello. What can I help you with?", body)]
    inside = [
        [Paragraph("Player choice", h2), Paragraph("Dialogue purpose", h2), Paragraph("Continuation", h2)],
        [Paragraph("A · There are far more kinds of books here than I imagined. / It really does look as if the rain has started outside.", body), Paragraph("Ordinary conversation or a connection to information about another resident", body), Paragraph("Maya responds with bookshop conversation, weather observations or a lead involving another NPC.", body)],
        [Paragraph("B · Do you have the newest edition?", body), Paragraph("A relationship-specific interaction", body), Paragraph("The exchange can lead to a recommendation or a discussion of the book's subject.", body)],
        [Paragraph("C · I'm here to discuss something serious.", body), Paragraph("Task recognition", body), Paragraph("Maya recognises the purpose of the visit and continues into the relevant task text.", body)],
    ]
    table2 = Table(inside, colWidths=[95*mm, 82*mm, 86*mm], repeatRows=1)
    table2.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),HexColor("#e6543f")),("TEXTCOLOR",(0,0),(-1,0),HexColor("#ffffff")),("VALIGN",(0,0),(-1,-1),"TOP"),("GRID",(0,0),(-1,-1),0.5,HexColor("#d5d6d2")),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
    story += [table2, Spacer(1, 8*mm), Paragraph("Follow-up contact", h1), Paragraph("The same three-part structure continues after the initial encounter: A can speak with Maya, B can enter the shop, and the recognition branch can progress. The content must change according to the protagonist's knowledge, prior choices and relationship with Maya rather than repeating one universal script.", body)]
    doc.build(story)


build_text_design_pdf()


make_doc(
    "week1-mist-harbor-story-outline-en.docx",
    "Mist Harbor Story Outline",
    "熊韬炀",
    "A complete narrative framework for Mist Harbor: a small coastal town shaped by an abandoned railway, a reduced port and one hundred permanent residents whose daily routines form the playable social world.",
    [
        ("Town and residency", [
            ("p", "Mist Harbor lies between mountain and sea at the end of an old railway and beside a small port. Around three thousand people live there, while the game follows roughly one hundred permanent residents with stable routines, private histories and ordinary traces of life. Applicants have seven days to submit basic documentation, record their presence in local life and receive recognition from at least twelve permanent residents. Recognition means only that someone is willing to say, 'I know this applicant.' It cannot be bought, and every resident keeps the right to refuse."),
        ]),
        ("History and atmosphere", [
            ("p", "Mist Harbor began as a market between fishing communities and people from the mountains. Railway and port expansion turned it into a transfer point for timber, herbs, stone, fish, salt and dried goods. Hotels, cafés, warehouses and a small theatre grew around the traffic. Forty years ago, a new main line bypassed the town and larger ports absorbed its freight. Services dwindled, cranes rusted and younger residents left."),
            ("p", "The town did not disappear. Remaining residents rebuilt life at a smaller, more predictable scale. Warehouses became a library, craft shops and a theatre. The old canteen kept recipes from the busiest years. Trains and boats still arrive, only more slowly. Daily schedules now repeat with unusual clarity: the same people walk dogs, study, work, read, play sport or watch sunset at recognisable hours."),
            ("p", "Residents sometimes say that the last freight train carried away part of the town's time. Sunset lingers, old light overlaps with the present in photographs, and rain occasionally contains the sound of wheels or doors that no longer exist. Nobody proves or explains these events. They sit beside the railway, tides and routines as part of ordinary life."),
        ]),
        ("Seven-day structure", [
            ("bullet", "Days 1–3 are light, playful and slightly absurd. A follows encounters and changing relationship chains. B plans around opening hours, transport and complete blocks of time."),
            ("bullet", "Days 4–5 give recognition weight. Residents may be unavailable, refuse because of an earlier encounter or offer recognition after an unexpected shared event."),
            ("bullet", "Day 6 forces decisions about the last uninterrupted blocks of time. An overlooked person may become central; an apparently secure recognition may change."),
            ("bullet", "Day 7 reviews the material, relationships and traces left in town. The ending is not a simple success/failure judgement. It reflects how each protagonist lived and was remembered."),
            ("p", "Short, lightly interactive transitions use photographs, drawing, collage and sound to connect the two simultaneous lives without turning their emotions into numbers."),
        ]),
        ("Activities and systems", [
            ("p", "Restaurant work and the public recipe book, commissioned writing and collage poetry, card reasoning and tarot, dialogue for resolving misunderstandings, three-dimensional optical illusion spaces, field recording and music production, chess, and moments of observation all grow from NPC routines rather than appearing as a compulsory checklist."),
            ("p", "Time, money, information and relationships shape how the activities are reached. A often buys speed or accepts an invitation before knowing the outcome. B protects money through planning and uses her notebook to preserve time. Both can reach the end, but the same opportunity carries a different cost. The recognition system therefore advances the story and the player's understanding of the town rather than functioning as a detached score."),
        ]),
    ],
)


make_doc(
    "week1-mist-harbor-story-revision-en.docx",
    "Mist Harbor Setting Revision",
    "熊韬炀",
    "This revision gives Mist Harbor a more specific coastal identity: limestone hills, lavender fields, an old railway and a small port shaped by the light, colour and daily rhythms of southern France.",
    [
        ("Landscape and architecture", [
            ("p", "Mist Harbor sits between a slope facing the sea and limestone hills warmed to pale gold and ochre. Olive trees, old vines and fields of lavender move in the wind. The air carries thyme, sea salt, dry grass and sun-heated stone. Most houses are pale terracotta, lime white or natural ochre, with faded blue or grey-green shutters. Cobbled lanes include shallow drainage channels, and summer brings fallen flowers and continuous cicadas."),
            ("p", "A mature plane tree and an ancient stone fountain occupy the central square. Residents move wicker chairs into the shade in the morning and evening. Mist lifts from the pebble shore as the sun rises, and sunset turns stone walls from yellow to orange and pink until they seem to merge with the pale cliffs beyond."),
        ]),
        ("Growth, decline and survival", [
            ("p", "The settlement began as a small port where fishers and mountain communities exchanged goods. Roman stonework and a damaged fountain remain. In the Middle Ages it linked coast and interior. Later, railway construction and port expansion brought ships carrying olive oil, wine, timber and stone. Railway workers, sailors, olive pressers and seasonal lavender and grape workers became the ancestors of many present residents."),
            ("p", "A motorway and a larger port bypassed the town roughly forty years ago. Freight stopped, railway sleepers filled with weeds and wildflowers, and shutters remained closed for years. Yet the town survived by contracting. Warehouses became a library, workshops and a theatre. The port serves fishing and small sailing boats. Markets open twice a week. Lavender fields and olive groves remain because they belong to the place, not because they still support mass export."),
        ]),
        ("Residency and time", [
            ("p", "Fifteen years ago, residents rejected the idea of solving depopulation through aggressive incentives. Instead, applicants must spend seven days inside the town's ordinary life and receive twelve statements of recognition. Affection is not required. Presence is. Refusal remains possible. The rule asks applicants to enter other people's rhythms with their own time and bodies: to be present in the square during the loudest cicadas, to remain beside an ochre wall at sunset, or to help pass a crate of tomatoes or lavender when the market closes."),
            ("p", "Local stories say the final freight train carried away part of the town's time. Summer sunsets linger, old and present light overlap on stone, and rain briefly returns the smell of steam or the sound of shutters as they were decades earlier. These details are never explained. They appear when a character's emotions loosen enough to notice them."),
        ]),
        ("Narrative use", [
            ("p", "The seven-day arc remains light and exploratory at first, turns towards the difficulty of recognition in its middle and lets earlier choices acquire weight during the final two days. Both protagonists meet the same people and spaces, but their time, money, information and willingness to accept chance produce different encounters. Activities grow from the town's work, craft, music and public life, so play and story remain part of one rhythm."),
        ]),
    ],
)


make_doc(
    "week1-solmere-story-outline-en.docx",
    "Solmere Story Outline",
    "HuieChen 陈慧娥",
    "Solmere is a coastal island town of around three thousand people. One hundred residents with regular schedules form its playable community, while recognition from twelve of them becomes the defining condition of a seven-day trial residency.",
    [
        ("The town", [
            ("p", "Solmere lies between mountain and sea. Its long summer light, pale stone, blue-green shutters, awnings, vines and lemon trees recall Provence and southern Italy. Yet it is a working town rather than a resort. Handwritten menus, laundry, small production notices and studio posters share the same streets."),
            ("p", "A former railway and reduced port left warehouses and large spaces behind. Film crews first came for the coastline, old town, steady light and low-cost buildings. A production company later turned a warehouse into a studio, the local school expanded into a media college, and post-production, sound, advertising, animation, games and music followed. The result resembles Los Angeles at a small-town scale: people edit in cafés, the record-shop owner takes sound commissions, and an unmarked old building may be a working studio."),
            ("p", "The industry supports stable work without erasing ordinary life. Public housing and controlled development keep rent and travel within reach. Residents may enjoy the familiar social world or find it restrictive. Solmere offers comfort, but comfort means something different to each person."),
        ]),
        ("Trial residency", [
            ("p", "Administration is ordinary except for one condition: before seven days end, an applicant must receive recognition from twelve residents. Recognition cannot be purchased through property, companies or family connections. It means that another person accepts that the applicant was genuinely present here. Money and time still influence how a person travels, waits, misses an opportunity and becomes memorable, so the rule does not pretend to erase unequal starting points."),
            ("p", "Solmere's time carries a slight magical-realist distortion. Older residents joke that the last freight train carried away a small part of the town's time. Evening light lingers. Old photographs overlap with the present. Field recordings sometimes contain wheels from trains that no longer run. The game does not explain these events; they surface briefly at emotional thresholds."),
        ]),
        ("Parallel seven-day arc", [
            ("bullet", "Day 1 introduces the rule and two ways of entering the town. A follows people and chance. B records opening hours, routes and NPC schedules."),
            ("bullet", "Days 2–3 grow relationships through activities. The protagonists meet some of the same people but hear different things because the timing, method and state of the encounter differ."),
            ("bullet", "Days 4–5 make recognition harder. Availability, earlier impressions and unexpected shared events replace the feeling of a simple collection task."),
            ("bullet", "Day 6 gives the final uninterrupted hours real weight. A notices how quickly she turns people into a project to optimise. B has to distinguish between what she can handle and what she should accept."),
            ("bullet", "Day 7 leaves each woman with a small action rather than a narrated lesson. Two near-identical sunset photographs close the parallel routes."),
        ]),
        ("A · Creative Direction", [
            ("p", "A is twenty-five, ambitious, optimistic and skilled at finding direction inside a confused project. Her family gave her time, cultural access and room to fail. She knows that support matters and sometimes discounts her own ability because of it. An early professional opportunity reached her through a family introduction, even though the work itself succeeded on its own quality. She comes to Solmere because its creative ecosystem attracts her and because the residency cannot simply be completed by her family. Her route asks her to accept that privilege and genuine ability can both be true."),
            ("p", "A acts before refining the plan. She accepts invitations, enters interesting places, cooks, watches people and can spend an afternoon without a measurable result. Creative direction also makes her sensitive to authority: she must decide what to keep, what to remove and how to hold a team together without turning her own judgement into the only answer."),
        ]),
        ("B · Production Management", [
            ("p", "B is twenty-five and grew up in a stable working household where money and opportunities had to be counted. She learned to check belongings, files, routes and consequences because avoidable mistakes carried real cost. In production management she excels at budgets, schedules, staffing, locations, equipment, delivery and crisis response. She wants larger projects and greater responsibility, but reliability makes her the person who absorbs everyone else's unfinished work."),
            ("p", "B comes to Solmere for a sustainable career in production, post-production and technical work, with living costs that make a long-term life possible. She is not trying to become spontaneous. Her route asks her to separate 'I can do this' from 'I must take this on.' Her notebook and sound work hold magical traces of former addresses, moving labels and earlier schedules before returning to records written in her own hand."),
        ]),
        ("Shared design principle", [
            ("p", "A and B are equally ambitious and hopeful. They do not exist to teach or rescue one another. The same residents, workplace and rules carry different weight because A begins with more room to recover while B begins by calculating consequences. Photographs, notebooks, sound, chess, cooking, collage and small acts of observation carry the hidden emotional line. The visible line remains clear: live in Solmere for seven days and become known by twelve people."),
        ]),
    ],
)
